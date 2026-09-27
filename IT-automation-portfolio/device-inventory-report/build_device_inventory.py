import openpyxl
from datetime import datetime
from collections import Counter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

src = openpyxl.load_workbook('New Hire Tracker.xlsx', data_only=True)['Inventory']
today = datetime.today()
headers = [src.cell(1,c).value for c in range(1,12)]
rows = []
for r in range(2, src.max_row+1):
    if src.cell(r,3).value in (None,''): continue
    v = [src.cell(r,c).value for c in range(1,12)]
    if isinstance(v[9], datetime): v[10] = (today - v[9]).days   # recompute idle
    rows.append(v)
order = {'Recover - Leaver':0,'Spare - Unassigned':1,'Idle 30d+':2,'Unknown User':3,'Assigned':4}
rows.sort(key=lambda v:(order.get(v[0],9), v[1] or '', v[2] or ''))

wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'Summary'
ws['A1'] = 'Device Inventory'; ws['A1'].font = Font(size=16, bold=True)
cs = Counter(v[0] for v in rows); ct = Counter(v[1] for v in rows)
ws['A3'] = 'By Status'; ws['A3'].font = Font(bold=True, size=12)
r = 4
for name, col in [('Recover - Leaver','C00000'),('Spare - Unassigned','00B050'),
                  ('Idle 30d+','FFC000'),('Unknown User','BFBFBF'),('Assigned','404040')]:
    if cs.get(name):
        ws.cell(r,1,name).font = Font(bold=True, color='FFFFFF')
        ws.cell(r,1).fill = PatternFill('solid', fgColor=col)
        ws.cell(r,2,cs[name]); r += 1
ws.cell(r,1,'TOTAL').font = Font(bold=True); ws.cell(r,2,len(rows)).font = Font(bold=True)

wi = wb.create_sheet('Inventory')
for c,h in enumerate(headers,1):
    cell = wi.cell(1,c,h); cell.fill = PatternFill('solid', fgColor='1F3864'); cell.font = Font(bold=True, color='FFFFFF')
border = Border(bottom=Side(style='thin', color='D9D9D9'))
for ri,v in enumerate(rows,2):
    for c,val in enumerate(v,1):
        cell = wi.cell(ri,c,val)
        if c in (9,10) and isinstance(val, datetime): cell.number_format = 'm/d/yy'
        if c in (5,7): cell.number_format = '@'
        cell.border = border
wi.freeze_panes = 'A2'; last = len(rows)+1; wi.auto_filter.ref = f'A1:K{last}'
def fill(h): return PatternFill(start_color=h, end_color=h, fill_type='solid')
for txt,f,w in [('Recover - Leaver',fill('C00000'),True),('Spare - Unassigned',fill('00B050'),False),
                ('Idle 30d+',fill('FFC000'),False),('Unknown User',fill('BFBFBF'),False)]:
    kw = {'fill':f}
    if w: kw['font'] = Font(color='FFFFFF')
    wi.conditional_formatting.add(f'A2:A{last}', FormulaRule(formula=[f'$A2="{txt}"'], **kw))
for c,wd in enumerate([18,8,26,42,16,22,34,11,11,13,10],1):
    wi.column_dimensions[get_column_letter(c)].width = wd
wb.save('Device Inventory.xlsx')
